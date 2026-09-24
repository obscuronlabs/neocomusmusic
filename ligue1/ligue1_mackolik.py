#!/usr/bin/env python3
"""Fransa Ligue 1 N. HAFTA.xlsx doldurucu — tek kaynak: arsiv.mackolik.com JSON.

Kullanım:
    python ligue1_mackolik.py 1 [--dir KLASOR] [--template SABLON.xlsx]

D–AC  = WeeklyStandingData hft=N (API sırası aynen, sort yok)
AL–BD = FixtureHandler week=N+1
Formüllere (AF-AK = 32..37 ve >=57) asla yazılmaz.
Denetim FAIL verirse çıktı "... HAFTA.FAILED.xlsx" olarak kaydedilir, asıl dosya yazılmaz.
"""
import argparse
import glob
import json
import os
import re
import shutil
import sys
import time
import urllib.parse
import urllib.request

import openpyxl
from openpyxl.utils import get_column_letter

BASE = "https://arsiv.mackolik.com"
HEADERS = {
    "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
                  "(KHTML, like Gecko) Chrome/124.0 Safari/537.36",
    "X-Requested-With": "XMLHttpRequest",
    "Referer": BASE + "/Standings/Default.aspx",
}
SLEEP = 0.4

# Sütunlar (1-based)
C_RANK, C_TEAM = 2, 4
C_O, C_G, C_B, C_M, C_A, C_Y, C_P = 5, 6, 7, 8, 9, 10, 11
C_AV = 13
C_HOME0, C_AWAY0 = 14, 22  # N–U, V–AC (8'er sütun)
C_SEASON, C_WEEK = 30, 31
C_AF = 32
C_DATE, C_MS, C_HRANK, C_HOME, C_SCORE, C_AWAY, C_ARANK, C_HT = 38, 39, 40, 41, 43, 45, 46, 47
C_ODDS0 = 49  # AW..BD (8)
PROTECTED = set(range(32, 38))  # AF..AK
FIRST_FORMULA_TAIL = 57         # BE ve sonrası

STANDING_COLS = [C_RANK, C_TEAM, C_O, C_G, C_B, C_M, C_A, C_Y, C_P, C_AV] + \
    list(range(C_HOME0, C_HOME0 + 8)) + list(range(C_AWAY0, C_AWAY0 + 8)) + [C_SEASON, C_WEEK]
MATCH_COLS = [C_DATE, C_MS, C_HRANK, C_HOME, C_SCORE, C_AWAY, C_ARANK, C_HT] + \
    list(range(C_ODDS0, C_ODDS0 + 8))
DEFAULT_SPLIT_ORDER = ["O", "G", "B", "M", "A", "Y", "P", "AV"]

LOG_LINES = []


def log(msg):
    print(msg, flush=True)
    LOG_LINES.append(msg)


# ---------------------------------------------------------------- HTTP / JSON
def _parse(text):
    text = text.strip()
    if not text:
        return None
    try:
        return json.loads(text)
    except ValueError:
        pass
    try:
        return json.loads(text.replace("'", '"'))
    except ValueError:
        return None


def get_json(path, params):
    url = BASE + path + "?" + urllib.parse.urlencode(params, safe="/")
    time.sleep(SLEEP)
    req = urllib.request.Request(url, headers=HEADERS)
    try:
        with urllib.request.urlopen(req, timeout=30) as r:
            raw = r.read().decode("utf-8", errors="replace")
    except Exception as e:  # HTTP 500 dahil -> blok boş
        log(f"  HTTP HATA {url} -> {e}")
        return None
    data = _parse(raw)
    if data is None:
        log(f"  JSON PARSE HATA {url}")
    return data


def _rows(data, key):
    if isinstance(data, dict):
        v = data.get(key)
        return v if isinstance(v, list) else []
    if isinstance(data, list):
        return data
    return []


# ---------------------------------------------------------------- İsimler
def canonical(name):
    n = str(name or "").strip()
    low = n.lower()
    if low in ("psg", "paris sg", "paris saint germain", "paris saint-germain") or "paris s" in low:
        return "Paris Saint-Germain"
    if "strasbourg" in low or "strousburg" in low:
        return "RC Strasbourg"
    if "lyon" in low:
        return "Lyon"
    if "monaco" in low:
        return "Monaco"
    if "marsilya" in low or "marseille" in low:
        return "Marsilya"
    return n


def to_int(x):
    try:
        return int(str(x).strip())
    except (TypeError, ValueError):
        return None


# ---------------------------------------------------------------- Mackolik
_SID_CACHE = {}


def season_id(start_year):
    if start_year in _SID_CACHE:
        return _SID_CACHE[start_year]
    data = get_json("/AjaxHandlers/CompetitionHandler.aspx",
                    {"op": "seasons", "group": 2, "year": f"{start_year}/{start_year + 1}"})
    sid = None
    for item in _rows(data, "s") if isinstance(data, dict) else (data or []):
        if isinstance(item, (list, tuple)) and len(item) > 1 and \
                str(item[1]).strip().lower() in ("ligue 1", "1.lig"):
            sid = item[0]
            break
    _SID_CACHE[start_year] = sid
    return sid


def name_map(sid):
    data = get_json("/AjaxHandlers/StandingHandler.ashx", {"op": "standing", "id": sid})
    m = {}
    for r in _rows(data, "s"):
        if isinstance(r, (list, tuple)) and len(r) > 1:
            m[str(r[0])] = canonical(r[1])
    return m


def weekly_standing(sid, week, names):
    data = get_json("/Standings/Data/WeeklyStandingData.aspx", {"seas": sid, "hft": week})
    out = []
    for t in _rows(data, "s"):  # API sırası = tablo sırası; SORT YOK
        if not isinstance(t, (list, tuple)) or len(t) < 22:
            continue
        tid = str(t[0])
        name = names.get(tid) or (canonical(t[1]) if isinstance(t[1], str) else None)
        v = [to_int(x) for x in t]
        hp, ap, hw, aw, hd, ad, hl, al = v[2], v[3], v[4], v[5], v[6], v[7], v[8], v[9]
        hgf, agf, hga, aga = v[10], v[11], v[12], v[13]
        hpts, apts = v[14], v[15]
        tot_p = (v[21] or 0) + (v[17] or 0)
        row = dict(id=tid, team=name,
                   O=hp + ap, G=hw + aw, B=hd + ad, M=hl + al,
                   A=hgf + agf, Y=hga + aga, P=tot_p,
                   home=dict(O=hp, G=hw, B=hd, M=hl, A=hgf, Y=hga, P=hpts, AV=hgf - hga),
                   away=dict(O=ap, G=aw, B=ad, M=al, A=agf, Y=aga, P=apts, AV=agf - aga))
        row["AV"] = row["A"] - row["Y"]
        out.append(row)
    return out


def fixtures(sid, week, names):
    data = get_json("/AjaxHandlers/FixtureHandler.aspx",
                    {"command": "getMatches", "id": sid, "week": week})
    out = []
    for m in _rows(data, "m"):
        if not isinstance(m, (list, tuple)) or len(m) < 11:
            continue

        def team(i_name):
            tid = str(m[i_name - 1]) if i_name >= 1 else None
            if tid in names:
                return names[tid]
            return canonical(m[i_name])

        hs, as_ = str(m[9]).strip(), str(m[10]).strip()
        played = hs != "" and as_ != "" and to_int(hs) is not None and to_int(as_) is not None
        ht = str(m[-2]).strip() if len(m) >= 2 else ""
        odds = list(m[11:19])
        out.append(dict(date=m[1], home=team(4), away=team(6),
                        score=f"{hs} - {as_}" if played else None,
                        ht=ht if played and ht else None,
                        odds=odds))
    return out


# ---------------------------------------------------------------- Excel
def season_from_header(text):
    s = str(text)
    m = re.search(r"(\d{4})\s*/\s*(\d{4})", s)
    if m:
        return int(m.group(1))
    m = re.search(r"(\d{2})\s*/\s*(\d{2})", s)
    if m:
        yy = int(m.group(1))
        return 2000 + yy if yy < 90 else 1900 + yy
    return None


def find_blocks(ws):
    heads = [r for r in range(1, ws.max_row + 1)
             if isinstance(ws.cell(r, C_RANK).value, str) and "HAFTA" in ws.cell(r, C_RANK).value.upper()]
    blocks = []
    for i, h in enumerate(heads):
        end = (heads[i + 1] - 1) if i + 1 < len(heads) else ws.max_row
        tk = None
        for r in range(h + 1, end + 1):
            v = ws.cell(r, C_TEAM).value
            if isinstance(v, str) and v.strip().upper() == "TAKIMLAR":
                tk = r
                break
        if tk is None:
            continue
        blocks.append(dict(head=h, takim=tk, first=tk + 1, last=end,
                           year=season_from_header(ws.cell(h, C_RANK).value)))
    return blocks


def split_order(ws, header_row, c0):
    labels = [str(ws.cell(header_row, c).value or "").strip().upper() for c in range(c0, c0 + 8)]
    if sorted(labels) == sorted(DEFAULT_SPLIT_ORDER):
        return labels
    return DEFAULT_SPLIT_ORDER


def _merged_at(ws, r, c):
    for rng in ws.merged_cells.ranges:
        if rng.min_row <= r <= rng.max_row and rng.min_col <= c <= rng.max_col:
            return rng
    return None


def safe_write(ws, r, c, value):
    if c in PROTECTED or c >= FIRST_FORMULA_TAIL:
        return
    cur = ws.cell(r, c).value
    if isinstance(cur, str) and cur.startswith("="):
        return  # formül üzerine asla yazma
    rng = _merged_at(ws, r, c)
    if rng is not None:
        ws.unmerge_cells(str(rng))
    ws.cell(r, c).value = value


def clear_rows(ws, rows, cols):
    for r in rows:
        for c in cols:
            safe_write(ws, r, c, None)


# ---------------------------------------------------------------- Ana akış
def process_block(ws, b, week, fails):
    year = b["year"]
    tag = f"{year}/{str(year + 1)[-2:]}" if year else f"satır {b['head']}"
    data_rows = list(range(b["first"], b["last"] + 1))
    af_before = {r: ws.cell(r, C_AF).value for r in data_rows}

    def fail(msg, row=None):
        fails.append(f"{tag}{'' if row is None else f' satır {row}'}: {msg}")

    if year is None:
        fail("başlıkta yıl yok")
        return
    max_week = 34 if year >= 2023 else 38
    if week > max_week:
        clear_rows(ws, data_rows, STANDING_COLS + MATCH_COLS)
        log(f"{tag}: 18 takımlı format, hafta {week} > {max_week} -> blok boş")
        return

    sid = season_id(year)
    if sid is None:
        clear_rows(ws, data_rows, STANDING_COLS + MATCH_COLS)
        if year >= 2026:
            log(f"{tag}: sId yok -> blok boş")
        else:
            fail("sId bulunamadı, blok boş")
        return
    names = name_map(sid)
    table = weekly_standing(sid, week, names)
    if not table:
        clear_rows(ws, data_rows, STANDING_COLS + MATCH_COLS)
        fail(f"WeeklyStandingData boş (sId={sid}, hft={week}), blok boş")
        return
    matches = [] if week >= max_week else fixtures(sid, week + 1, names)

    # Satır yetmezse bloğun sonuna ekle (ortada boşluk yok)
    need = len(table)
    if len(data_rows) < need:
        add = need - len(data_rows)
        ws.insert_rows(b["last"] + 1, add)
        b["last"] += add
        data_rows = list(range(b["first"], b["last"] + 1))
        af_before.update({r: None for r in data_rows if r not in af_before})
        log(f"{tag}: {add} satır eklendi")

    home_order = split_order(ws, b["takim"], C_HOME0)
    away_order = split_order(ws, b["takim"], C_AWAY0)
    rank_of = {}
    clear_rows(ws, data_rows, STANDING_COLS + MATCH_COLS)
    for i, t in enumerate(table):
        r = b["first"] + i
        rank_of[t["team"]] = i + 1
        safe_write(ws, r, C_RANK, i + 1)
        safe_write(ws, r, C_TEAM, t["team"])
        for c, k in zip((C_O, C_G, C_B, C_M, C_A, C_Y, C_P), "O G B M A Y P".split()):
            safe_write(ws, r, c, t[k])
        safe_write(ws, r, C_AV, t["AV"])
        for j, k in enumerate(home_order):
            safe_write(ws, r, C_HOME0 + j, t["home"][k])
        for j, k in enumerate(away_order):
            safe_write(ws, r, C_AWAY0 + j, t["away"][k])
        safe_write(ws, r, C_SEASON, f"{year}/{year + 1}")
        safe_write(ws, r, C_WEEK, week)

    for i, m in enumerate(matches):
        r = b["first"] + i
        if r > b["last"]:
            fail("maç sayısı blok satırlarını aşıyor", r)
            break
        safe_write(ws, r, C_DATE, m["date"])
        safe_write(ws, r, C_MS, "MS" if m["score"] else None)
        safe_write(ws, r, C_HRANK, rank_of.get(m["home"]))
        safe_write(ws, r, C_HOME, m["home"])
        safe_write(ws, r, C_SCORE, m["score"])
        safe_write(ws, r, C_AWAY, m["away"])
        safe_write(ws, r, C_ARANK, rank_of.get(m["away"]))
        safe_write(ws, r, C_HT, m["ht"])
        for j, o in enumerate(m["odds"][:8]):
            safe_write(ws, r, C_ODDS0 + j, o if o not in ("", None) else None)

    log(f"{tag}: sId {sid} … puan {len(table)} mac {len(matches)}")
    audit(ws, b, tag, week, max_week, table, matches, af_before, fail)


def audit(ws, b, tag, week, max_week, table, matches, af_before, fail):
    rows = list(range(b["first"], b["last"] + 1))
    teams = [ws.cell(r, C_TEAM).value for r in rows if ws.cell(r, C_TEAM).value not in (None, "")]
    # 1
    if len(teams) != len(table):
        fail(f"[1] D sayısı {len(teams)} != api.s {len(table)}")
    # 2
    r0 = b["first"]
    x = (ws.cell(r0, C_TEAM).value, ws.cell(r0, C_P).value, ws.cell(r0, C_AV).value)
    a = (table[0]["team"], table[0]["P"], table[0]["AV"])
    if x != a:
        fail(f"[2] 1. satır Excel {x} != API {a}", r0)
    # 3
    for i in range(len(table)):
        r = r0 + i
        g = lambda c: ws.cell(r, c).value or 0
        if g(C_O) != g(C_G) + g(C_B) + g(C_M):
            fail("[3] O != G+B+M", r)
        if g(C_P) != 3 * g(C_G) + g(C_B):
            fail(f"[3] P != 3G+B ({g(C_P)} vs {3 * g(C_G) + g(C_B)})", r)
        if g(C_AV) != g(C_A) - g(C_Y):
            fail("[3] AV != A-Y", r)
    # 4
    if week < max_week:
        xl_pairs = [(ws.cell(r, C_HOME).value, ws.cell(r, C_AWAY).value)
                    for r in rows if ws.cell(r, C_HOME).value]
        api_pairs = [(m["home"], m["away"]) for m in matches]
        xl_ft = sum(1 for r in rows if ws.cell(r, C_SCORE).value)
        api_ft = sum(1 for m in matches if m["score"])
        if xl_pairs != api_pairs or xl_ft != api_ft:
            fail(f"[4] maç eşleşmesi/FT farklı (excel {len(xl_pairs)}/{xl_ft}, api {len(api_pairs)}/{api_ft})")
        if not matches:
            fail(f"[4] hafta {week + 1} fikstürü boş")
        # 5
        in_fix = {n for p in xl_pairs for n in p}
        dset = set(teams)
        for n in in_fix - dset:
            fail(f"[5] maç bloğundaki '{n}' D sütununda yok")
        byes = dset - in_fix
        if byes and len(byes) != len(dset) - 2 * len(xl_pairs):
            fail(f"[5] eşleşmeyen takımlar: {sorted(byes)}")
        elif byes:
            log(f"{tag}: bye/oynamayan: {sorted(byes)}")
    else:
        if any(ws.cell(r, C_HOME).value for r in rows):
            fail("[4] son hafta ama maç bloğu dolu")
    # 6
    for r, v in af_before.items():
        if isinstance(v, str) and v.startswith("="):
            now = ws.cell(r, C_AF).value
            if not (isinstance(now, str) and now.startswith("=")):
                fail("[6] AF formülü bozuldu", r)
    # 7
    if b["year"] >= 2023 and week > 34 and teams:
        fail("[7] 18 takımlı yılda hafta>34 ama blok dolu")


def resolve_paths(folder, week, template):
    out = os.path.join(folder, f"Fransa Ligue 1 {week}. HAFTA.xlsx")
    if os.path.exists(out):
        return out, out
    if template:
        return template, out
    cands = sorted(glob.glob(os.path.join(folder, "Fransa Ligue 1 *. HAFTA.xlsx")))
    if not cands:
        sys.exit(f"Şablon bulunamadı: {folder} içinde 'Fransa Ligue 1 N. HAFTA.xlsx' yok (--template ver)")
    return cands[0], out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("week", type=int)
    ap.add_argument("--dir", default=".")
    ap.add_argument("--template")
    args = ap.parse_args()
    N = args.week
    if not 1 <= N <= 38:
        sys.exit("hafta 1..38")

    src, out = resolve_paths(args.dir, N, args.template)
    if src != out:
        shutil.copyfile(src, out + ".tmp")
        src = out + ".tmp"
    wb = openpyxl.load_workbook(src)
    ws = wb.active
    blocks = find_blocks(ws)
    log(f"Dosya: {out}  sayfa: {ws.title}  blok: {len(blocks)}  hafta={N} (maç hafta={N + 1})")

    fails = []
    # Satır ekleme alttaki blokları kaydırdığı için sondan başa işle
    for b in sorted(blocks, key=lambda b: b["head"], reverse=True):
        process_block(ws, b, N, fails)

    if fails:
        bad = out.replace(".xlsx", ".FAILED.xlsx")
        wb.save(bad)
        log(f"\nFAIL ({len(fails)}):")
        for f in fails:
            log("  " + f)
        log(f"Teslim yok. İnceleme için: {bad}")
        rc = 1
    else:
        wb.save(out)
        log(f"\nDenetim temiz ({len(blocks)} blok). Kaydedildi: {out}")
        rc = 0
    if src.endswith(".tmp"):
        os.remove(src)
    with open(out.replace(".xlsx", ".log"), "w", encoding="utf-8") as fh:
        fh.write("\n".join(LOG_LINES) + "\n")
    sys.exit(rc)


if __name__ == "__main__":
    main()
